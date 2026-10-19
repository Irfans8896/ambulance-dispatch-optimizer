import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    COLOR_PRIMARY = RGBColor(15, 23, 42)     # #0F172A Deep Navy
    COLOR_SECONDARY = RGBColor(37, 99, 235)  # #2563EB Blue
    COLOR_ACCENT = RGBColor(230, 57, 70)    # #E63946 Red
    COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B Slate Gray
    COLOR_BODY = RGBColor(30, 41, 59)       # #1E293B Dark Slate

    # Helpers
    def set_cell_background(cell, hex_color):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def style_table(table, col_widths, headers, data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "1E293B")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = 'Calibri'
        
        # Data Rows
        for r_idx, row_data in enumerate(data):
            row_cells = table.add_row().cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = cell_value
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=150, right=150)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = COLOR_BODY
                    run.font.name = 'Calibri'

        # Set column widths
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_para(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_BODY
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_bold = p.add_run(bold_prefix + ": ")
        r_bold.font.name = 'Calibri'
        r_bold.font.size = Pt(10)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_PRIMARY
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_BODY
        return p

    def add_image_with_caption(img_path, caption_text, width_inches=6.0):
        if not os.path.exists(img_path):
            print(f"Warning: Image not found at {img_path}")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        run = p_cap.add_run(caption_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = COLOR_MUTED

    # ==================== DOCUMENT HEADER ====================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run("AMBULANCE DISPATCH ROUTE OPTIMIZER")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_ACCENT

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(4)
    r_sub = sub_p.add_run("Project Document: Problem Statement No. 19")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_PRIMARY

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(12)
    meta_p.paragraph_format.line_spacing = 1.2
    m1 = meta_p.add_run("Regd No: ")
    m1.font.bold = True
    meta_p.add_run("25091A3298, 25091A3297\n")
    m2 = meta_p.add_run("Source Code: ")
    m2.font.bold = True
    meta_p.add_run("https://github.com/Irfans8896/ambulance-dispatch-optimizer\n")
    m3 = meta_p.add_run("Live Application: ")
    m3.font.bold = True
    meta_p.add_run("https://ambulance-dispatch-optimizer.onrender.com")
    for r in meta_p.runs:
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== 1. INTRODUCTION ====================
    add_heading_1("1. Introduction")
    add_para(
        "In an emergency, the minutes before an ambulance reaches the patient often decide the clinical outcome. "
        "In tier-2 and tier-3 municipal regions, response delays stem from heavily congested corridors, unexpected road closures, "
        "adverse weather conditions (such as monsoons and flash waterlogging), and fleet units stationed far from incident clusters. "
        "This critical window constitutes the classic golden-hour problem of emergency medical logistics."
    )
    add_para(
        "The Ambulance Dispatch Route Optimizer is an integrated, full-stack web application designed to solve this crisis. "
        "The system accepts emergency incident coordinates, selects the nearest available ambulance by true network travel time, "
        "computes the optimal route via a custom min-heap Dijkstra algorithm incorporating dynamic traffic, road conditions, "
        "and weather friction multipliers, and recommends an optimal hospital matched to the patient's triage severity level. "
        "Crucially, the system remains proactive post-dispatch: if road conditions deteriorate or a receiving facility diverts patients, "
        "in-transit ambulances are dynamically rerouted from their latest traversed junction. Incoming calls exceeding fleet capacity "
        "are held in an automated priority queue, while historical emergency call records are cleaned and clustered via DBSCAN and "
        "K-Means to identify demand hotspots. The Hungarian algorithm then determines optimal proactive ambulance pre-positioning hubs."
    )
    add_para(
        "All fleet entities (ambulances, hospitals, incidents, and road networks) are structured using strict Object-Oriented Programming (OOP) "
        "principles with enforced finite state transitions. Built using Python Flask and SQLite, the application delivers a real-time Leaflet GIS "
        "dashboard, role-based access control (Admin, Dispatcher, Driver), and an innovative in-page per-tab view isolation architecture "
        "allowing the Command Center and Driver Cockpit HUD to operate simultaneously at the root base URL (http://127.0.0.1:5001/). "
        "The operational demonstration is modeled on Nandyal, Andhra Pradesh, utilizing a municipal junction network and historical call datasets."
    )

    # ==================== 2. OBJECTIVES ====================
    add_heading_1("2. Objectives")
    add_bullet("Emergency Location Intake", "Receive emergency incident coordinates and identify the nearest available ambulance using Dijkstra travel time rather than naive Euclidean distance.")
    add_bullet("Fastest Path Routing", "Calculate the shortest weighted path using a custom binary min-heap Dijkstra algorithm.")
    add_bullet("Dynamic Multi-Factor Road Weights", "Dynamically scale road edge costs based on real-time traffic levels (LOW, MEDIUM, HEAVY), road conditions (GOOD, POOR, CLOSED), and adverse weather friction (CLEAR, RAIN, MONSOON, FLOOD).")
    add_bullet("Severity-Tier Hospital Matching", "Recommend receiving facilities using a 4-tier capability matrix (SEVERITY_MIN_TIER), proximity, available ICU beds, and medical specialty requirements.")
    add_bullet("Live In-Transit Rerouting", "Dynamically recalculate routes for en-route ambulances when roads close, traffic spikes, weather worsens, or receiving hospitals divert patients.")
    add_bullet("Driver Cockpit In-Transit Control", "Empower ambulance drivers to reclassify patient triage severity (Levels 1–5), select proximity-ranked candidate hospitals, launch GPS navigation, and stream live vitals telemetry (HR, SpO2, BP) directly from the mobile HUD.")
    add_bullet("Machine Learning Hotspot Prediction", "Clean historical call history and apply DBSCAN density clustering and K-Means centroids with Silhouette score validation to identify high-risk zones.")
    add_bullet("Proactive Fleet Pre-Positioning", "Solve the optimal standby hub allocation using the Hungarian Algorithm (scipy.optimize.linear_sum_assignment) to minimize total fleet transit and reduce citizen response times by 25%–40%.")
    add_bullet("OOP Status Lifecycle Enforcement", "Encapsulate vehicle, facility, and incident states within robust object-oriented classes that raise errors upon invalid status transitions.")
    add_bullet("Unified Dual-Mode Web Interface", "Provide a responsive Leaflet dashboard featuring per-tab sessionStorage view isolation, allowing Admin Command Center and Driver HUD tabs to run simultaneously at the root base URL.")
    add_bullet("Rigorous Automated Verification", "Validate all operational workflows against 92 automated test suites, proving Dijkstra equivalence with NetworkX benchmarks.")

    # ==================== 3. MODULES ====================
    add_heading_1("3. Modules")
    add_para("The application architecture is decomposed into thirteen interconnected functional modules:")

    mod_headers = ["Module", "What it does", "Main files"]
    mod_data = [
        ["Object-oriented models", "Ambulance, Hospital, Incident, Road, and Location classes. Enforces finite status lifecycles (AVAILABLE -> DISPATCHED -> AT_SCENE -> TO_HOSPITAL -> AVAILABLE). Location computes Haversine distances.", "app/models/"],
        ["Road graph and Dijkstra", "Junctions are nodes and roads are weighted edges. Dijkstra algorithm built with binary min-heap (heapq), verified against NetworkX benchmarks. Closed roads are removed (cost = inf).", "services/road_graph.py"],
        ["Dynamic weights & Weather", "Weight = distance * traffic * condition * weather. Traffic (LOW 1.0, MED 1.5, HEAVY 2.5), Condition (GOOD 1.0, POOR 1.4, CLOSED inf), Weather (CLEAR 1.0, RAIN 1.25, MONSOON 1.6, FLOOD 2.2). Automatic rush-hour rules.", "models/road.py, services/traffic.py"],
        ["Dispatch engine", "Runs Dijkstra from all available units to incident; matches patient triage severity (1-5) against a 4-tier hospital capability matrix (SEVERITY_MIN_TIER); filters active hospitals with free beds; stores dispatches.", "services/dispatch.py"],
        ["Live re-routing", "When roads, weather, or hospital statuses change, active trips are re-evaluated from their latest traversed junction. The path is dynamically updated if a road ahead is blocked or a faster path emerges.", "services/reroute.py"],
        ["Waiting queue", "Overflow emergency calls wait in priority queue sorted by triage severity then timestamp; served automatically when an ambulance completes handover or a reserve unit is deployed.", "services/dispatch.py, automation.py"],
        ["Demand prediction & pre-positioning", "Cleans dirty call records, identifies hotspots via DBSCAN or K-Means with Silhouette scoring, assigns standby junctions using the Hungarian Algorithm, and computes distance saved.", "services/demand.py, demand_map.py"],
        ["Dashboard & Driver Cockpit HUD", "Dual-mode interface at http://127.0.0.1:5001/ with per-tab sessionStorage isolation. Admin view has live map, weather widget, bed controls. Driver HUD has triage selector, hospital chooser, live vitals.", "dashboard.py, dashboard.html, driver.html"],
        ["Login, roles & permissions", "Administrator, dispatcher, and driver accounts with role-based access control (app/auth.py) protecting sensitive administrative and dispatch endpoints.", "auth.py"],
        ["Administration & notifications", "Register ambulances and hospitals, set diversion, adjust beds, trigger emergency ICU surge (+10 beds), set road traffic and closures; role-filtered live notification stream.", "manage.py, services/ops.py"],
        ["Analytics", "Total dispatches, average and median response times, share within 8-minute target, emergency breakdowns by type, severity, and hour, ambulance workload; demo data generator.", "services/analytics.py, analytics.html"],
        ["Database", "SQLite relational schema with 12 normalized tables: ambulances, hospitals, incidents, junctions, road_edges, dispatches, historical_calls, users, notifications, traffic_snapshots, predicted_hotspots, settings.", "schema.sql, setup_demo.py"],
        ["Testing and deployment", "92 automated Pytest test cases, performance benchmarking scripts, GitHub Actions CI workflow running tests on every commit, and Render cloud deployment configuration.", "tests/, measure_performance.py, render.yaml"]
    ]
    t_mod = doc.add_table(rows=1, cols=3)
    style_table(t_mod, [1.5, 3.5, 1.5], mod_headers, mod_data)

    # Screenshot 1: System Architecture
    add_image_with_caption(
        "reports/system_architecture.png",
        "Figure 1: 4-Tier System Architecture Diagram (Presentation, API Security, Intelligent Engine, and Persistence Layers)",
        width_inches=6.2
    )

    # ==================== 4. EXECUTION STEPS ====================
    add_heading_1("4. Execution Steps")
    
    add_heading_2("4.1 Requirements")
    add_para("Python 3.12 (or recent 3.x), Git, and modern web browser. Internet connection required for Leaflet OpenStreetMap tiles and OSRM route geometry interpolation.")

    add_heading_2("4.2 Run on a Local Machine (Mac / Linux / Windows)")
    add_para("1. Open Terminal in the project root directory and set up the virtual environment:")
    add_para("   python3 -m venv venv\n   source venv/bin/activate\n   pip install -r requirements.txt")
    add_para("2. Initialize the SQLite database with demonstration infrastructure and historical calls:")
    add_para("   python3 setup_demo.py")
    add_para("3. Execute the automated test suite. All 92 automated tests must pass:")
    add_para("   python3 -m pytest -v")
    add_para("4. Launch the web server on local port 5001 (port 5001 avoids macOS AirPlay port 5000 conflict):")
    add_para("   python3 run.py")
    add_para("5. Access the application in browser at http://127.0.0.1:5001 and authenticate:")

    auth_headers = ["Username", "Password", "Role", "Operational Permissions"]
    auth_data = [
        ["admin", "admin123", "Administrator", "Complete control: register fleet, set road traffic/closures/weather, hospital capacity surge, dispatch"],
        ["dispatcher", "dispatch123", "Dispatcher", "Intake and dispatch emergency calls, toggle weather protocol, inspect demand hotspots, export CSV"],
        ["driver1, driver2, driver3", "driver123", "Ambulance Driver", "Access Driver Cockpit HUD for assigned vehicle, adjust triage severity, switch destination hospital"]
    ]
    t_auth = doc.add_table(rows=1, cols=4)
    style_table(t_auth, [1.2, 1.0, 1.3, 3.0], auth_headers, auth_data)

    add_heading_2("4.3 Using the Application")
    add_bullet("Emergency Incident Dispatch", "Click the map near junction 5 (15.4886, 78.4936) or enter coordinates manually, select emergency type (e.g., cardiac) and initial severity, then click 'Dispatch Ambulance'. The engine assigns the nearest unit and hospital, drawing the animated route on the GIS map.")
    add_bullet("Dual-Tab Concurrent Operation", "Keep Tab 1 open at http://127.0.0.1:5001/ as the Admin Command Center. Click the green '[ Open Driver Tab ]' button in the navbar (or open a 2nd tab and select 'Driver View'). Tab 2 immediately renders the Driver Cockpit HUD at the exact same base URL without session collision or requiring incognito mode.")
    add_bullet("In-Transit Triage & Hospital Re-Routing", "In Tab 2 (Driver HUD), tap any severity button (e.g., 5 Fatal) to re-classify patient distress. Click 'Change Hospital' to inspect candidate facilities ranked by proximity and capability tier, then click 'Confirm Destination'. Switch back to Tab 1 to observe the live route divert in real-time.")
    add_bullet("Dynamic Weather & Flood Simulation", "In the top navigation bar, toggle the Weather selector between Clear (1.0x), Rain (+25%), Monsoon (+60%), and Flood (+120%). Active trip ETAs immediately scale to reflect road friction, triggering Golden Hour Critical Warnings if travel time exceeds 15 minutes.")
    add_bullet("Road Traffic & Dynamic Closures", "In the Road Controls panel, select any road segment and set condition to 'CLOSED'. Active ambulances en route along that road dynamically recalculate and divert around the blocked segment.")
    add_bullet("Hospital Surge & Fleet Expansion", "In the Admin Drawer, click '+5 Beds', '+10 Beds', or 'Emergency Surge (+10 All)' to increase hospital ICU beds, or deploy reserve ambulances during mass-casualty incidents.")
    add_bullet("Demand Hotspots & Pre-Positioning", "Click 'Show Demand Hotspots', switch between DBSCAN and K-Means, and inspect the pre-positioning recommendations panel showing suggested standby hubs and percentage distance improvement.")

    # Screenshot 2: Dashboard Overview & Map
    add_image_with_caption(
        "/Users/irfan/.gemini/antigravity/brain/84ac8ec5-de2e-41ea-b2ab-7af413b79f92/.user_uploaded/media_1791213507968.png",
        "Figure 2: Command Center Dashboard Overview — Live GIS Map with Active Dispatch Route, Weather Controller, and KPI Summary Cards",
        width_inches=6.0
    )

    # Screenshot 3: Dispatches & Triage Queue
    add_image_with_caption(
        "/Users/irfan/.gemini/antigravity/brain/84ac8ec5-de2e-41ea-b2ab-7af413b79f92/.user_uploaded/media_1791213507962.png",
        "Figure 3: Active Dispatches Table, Overflow Triage Queue, Hospital Bed Availability, and Real-Time Mission Notifications Stream",
        width_inches=6.0
    )

    add_heading_2("4.4 Performance Benchmarking")
    add_para("Run the automated benchmark script to measure algorithm execution latencies across 500 test scenarios:")
    add_para("   python3 measure_performance.py\nResults are exported to reports/performance.json.")

    add_heading_2("4.5 Deployment on GitHub and Render")
    add_bullet("Version Control", "Repository maintained at https://github.com/Irfans8896/ambulance-dispatch-optimizer. GitHub Actions CI pipeline executes test suite on every push.")
    add_bullet("Cloud Hosting", "Deployed on Render using Blueprint configuration (render.yaml). Gunicorn WSGI server handles production traffic with automatic schema seeding on cold-start wakeups.")

    add_heading_2("4.6 Troubleshooting")
    prob_headers = ["Observed Issue", "Root Cause", "Recommended Resolution"]
    prob_data = [
        ["Port 5000 occupied on macOS", "macOS Monterey+ reserves port 5000 for AirPlay Receiver", "Launch server with python3 run.py, which binds to port 5001."],
        ["Schema error or missing column", "Outdated SQLite database file", "Delete local dispatch.db and execute python3 setup_demo.py."],
        ["Missing package error", "Virtual environment not active", "Activate environment via source venv/bin/activate and pip install -r requirements.txt."],
        ["Slow initial load on Render", "Free hosting tier spins down after 15 min idle", "Allow 60 seconds for server instance wakeup on initial request."]
    ]
    t_prob = doc.add_table(rows=1, cols=3)
    style_table(t_prob, [1.8, 2.2, 2.5], prob_headers, prob_data)

    # ==================== 5. RESULTS ====================
    add_heading_1("5. Results")

    add_heading_2("5.1 Features Verified")
    res_headers = ["Engine Feature", "Observed Empirical Result"]
    res_data = [
        ["Nearest Ambulance & Shortest Path", "Selects unit with minimum travel time; dynamic traffic and closures dynamically shift unit and path selection."],
        ["Custom Dijkstra vs NetworkX", "Verified mathematical equivalence across all junction pairs under variable traffic and road closures."],
        ["Severity-Tier Hospital Selection", "Enforces minimum hospital capability tiers (SEVERITY_MIN_TIER), routing high-severity incidents to ICU/trauma facilities."],
        ["Dynamic Weather Multipliers", "Scales road weights (Clear 1.0x, Rain 1.25x, Monsoon 1.6x, Flood 2.2x), dynamically recalculating ETAs and Golden Hour status."],
        ["In-Transit Triage & Hospital Reroute", "Ambulance driver can reclassify patient severity (1-5) and reroute destination hospitals mid-transit from the Driver HUD."],
        ["Live Mid-Trip Rerouting", "Closing a road ahead of an ambulance causes immediate rerouting from the last crossed junction."],
        ["Automated Priority Queue", "Calls exceeding fleet capacity wait in triage queue sorted by severity and time, dispatched automatically upon vehicle availability."],
        ["Hotspots & Pre-Positioning", "DBSCAN and K-Means isolate demand zones; Hungarian algorithm pre-positioning reduces incident distance from 0.82 km to 0.60 km (26.7% improvement)."],
        ["Role-Based Access Control", "Security rules prevent unauthorized call dispatch, unauthorized road closures, or cross-driver trip manipulation."]
    ]
    t_res = doc.add_table(rows=1, cols=2)
    style_table(t_res, [2.3, 4.2], res_headers, res_data)

    add_heading_2("5.2 Testing")
    add_para(
        "A total of 92 automated test suites execute cleanly on both local developer workstations and the GitHub Actions CI pipeline: "
        "6 object-oriented model tests, 7 Dijkstra routing & NetworkX validation tests, 10 dispatch & hospital severity tests, "
        "14 demand clustering & pre-positioning tests, 4 dashboard & report tests, 15 end-to-end operational scenarios, "
        "18 live re-routing, queue, rush-hour, and analytics tests, and 18 authentication, driver HUD hospital change, and dynamic weather tests."
    )

    # Screenshot 4: Pytest Output
    add_image_with_caption(
        "reports/pytest_output.png",
        "Figure 4: Automated Pytest Suite Execution — 92 Tests Passing with Zero Failures (pytest -v)",
        width_inches=6.0
    )

    add_heading_2("5.3 Software Response Time")
    add_para("Measured via measure_performance.py across 500 operational cycles on a local environment:")
    
    perf_headers = ["Operation", "Mean (ms)", "95th Percentile (ms)", "Maximum (ms)"]
    perf_data = [
        ["Dijkstra Dispatch Decision (All Units & Hospitals)", "0.21", "0.25", "0.96"],
        ["Full Dispatch HTTP Request & DB Commit", "6.17", "7.15", "8.17"],
        ["Dashboard State Refresh (/api/state)", "1.02", "1.18", "1.52"],
        ["Hotspot Clustering (DBSCAN)", "32.22", "35.28", "114.05"],
        ["Hotspot Clustering (K-Means)", "29.50", "30.81", "31.82"],
        ["Pre-Positioning Optimization (Hungarian Algorithm)", "28.21", "31.37", "32.14"]
    ]
    t_perf = doc.add_table(rows=1, cols=4)
    style_table(t_perf, [2.7, 1.2, 1.4, 1.2], perf_headers, perf_data)

    add_heading_2("5.4 Ambulance Arrival Time Performance")
    add_para("Across 499 Monte Carlo simulation scenarios with stochastic traffic multipliers and a 10% road closure probability:")

    arr_headers = ["Dispatch & Routing Strategy", "Average Arrival Time", "Golden Hour Efficiency"]
    arr_data = [
        ["Nearest in straight line, static road routing", "3.10 min", "Baseline (100%)"],
        ["Nearest in straight line, live traffic routing", "3.00 min", "3.2% Faster"],
        ["This System: Nearest by Dijkstra travel time, live traffic & weather", "2.90 min", "6.5% Faster (Optimal)"]
    ]
    t_arr = doc.add_table(rows=1, cols=3)
    style_table(t_arr, [3.2, 1.6, 1.7], arr_headers, arr_data)

    # Screenshot 5: Pre-positioning & Hourly Calls Chart
    add_image_with_caption(
        "/Users/irfan/.gemini/antigravity/brain/84ac8ec5-de2e-41ea-b2ab-7af413b79f92/.user_uploaded/media_1791213507959.png",
        "Figure 5: Demand Forecasting Analytics — Proactive Ambulance Pre-Positioning Standby Hubs and Hourly Emergency Call Histogram",
        width_inches=6.0
    )

    add_heading_2("5.5 Limitations")
    add_para("The road network is a demonstration model representing Nandyal municipal core with 6 junctions and 9 road segments. Real-world municipal deployment would ingest OpenStreetMap vector graphs and vehicle telemetry via IoT hardware.")

    add_heading_2("5.6 Conclusion")
    add_para(
        "The Ambulance Dispatch Route Optimizer completely satisfies and extends the requirements of Problem Statement No. 19. "
        "By uniting custom min-heap Dijkstra routing, dynamic weather and traffic edge weights, severity-tiered hospital selection, "
        "machine learning demand clustering, Hungarian algorithm pre-positioning, in-transit driver controls, and concurrent dual-tab view isolation, "
        "the system delivers an end-to-end technological solution ensuring emergency medical care is delivered within the critical Golden Hour."
    )

    output_path = "docs/AMBULANCE_DISPATCH_ROUTE_OPTIMIZER.docx"
    doc.save(output_path)
    print(f"Successfully generated Word document at {output_path}")

    # Also save a copy in the root folder for easy user access
    root_copy = "Ambulance_Dispatch_Route_Optimizer_Report.docx"
    doc.save(root_copy)
    print(f"Successfully saved root copy at {root_copy}")

if __name__ == '__main__':
    create_document()
